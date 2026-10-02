"""Tests for core/project.py — text extraction and frontmatter parsing."""
import pytest
from pathlib import Path
from core.project import _extract_text, _parse_md_frontmatter


# ---------------------------------------------------------------------------
# _parse_md_frontmatter
# ---------------------------------------------------------------------------

def test_frontmatter_title(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\ntitle: Min anteckning\ncategory: Intervju\n---\n\nText här.", encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm["title"] == "Min anteckning"
    assert fm["category"] == "Intervju"


def test_frontmatter_quoted_values(tmp_path):
    f = tmp_path / "note.md"
    f.write_text('---\ntitle: "Med citattecken"\n---\nText.', encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm["title"] == "Med citattecken"


def test_frontmatter_single_quoted(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\ntitle: 'Enkelfnuttar'\n---\nText.", encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm["title"] == "Enkelfnuttar"


def test_frontmatter_tags_flow_style(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\ntags: [pilot, lärare, fas2]\n---\nText.", encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm["tags"] == ["pilot", "lärare", "fas2"]


def test_frontmatter_tags_block_style(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\ntags:\n  - pilot\n  - lärare\n  - fas2\n---\nText.", encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm["tags"] == ["pilot", "lärare", "fas2"]


def test_frontmatter_tags_quoted_values(tmp_path):
    f = tmp_path / "note.md"
    f.write_text('---\ntags: ["med mellanslag", \'enkelfnutt\']\n---\nText.', encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm["tags"] == ["med mellanslag", "enkelfnutt"]


def test_frontmatter_other_list_fields_ignored(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\nphotos:\n  - bild.jpg\n---\nText.", encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert "photos" not in fm


def test_frontmatter_no_frontmatter(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("Bara text, ingen frontmatter.", encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm == {}


def test_frontmatter_empty_file(tmp_path):
    f = tmp_path / "empty.md"
    f.write_text("", encoding="utf-8")
    fm = _parse_md_frontmatter(f)
    assert fm == {}


# ---------------------------------------------------------------------------
# _extract_text — YAML frontmatter stripping
# ---------------------------------------------------------------------------

def test_md_frontmatter_stripped(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\ntitle: Test\ncategory: Foo\n---\n\nDetta är texten.", encoding="utf-8")
    text = _extract_text(f)
    assert "title" not in text
    assert "category" not in text
    assert "Detta är texten" in text


def test_md_no_frontmatter_untouched(tmp_path):
    f = tmp_path / "plain.md"
    f.write_text("# Rubrik\n\nBrödtext.", encoding="utf-8")
    text = _extract_text(f)
    assert "Rubrik" in text
    assert "Brödtext" in text


def test_md_markdown_syntax_stripped(tmp_path):
    f = tmp_path / "md.md"
    f.write_text("**Fet** och _kursiv_ text.", encoding="utf-8")
    text = _extract_text(f)
    assert "**" not in text
    assert "_" not in text
    assert "Fet" in text
    assert "kursiv" in text


# ---------------------------------------------------------------------------
# _extract_text — encoding fallback
# ---------------------------------------------------------------------------

def test_utf8_file(tmp_path):
    f = tmp_path / "utf8.txt"
    f.write_text("Åäö är svenska tecken.", encoding="utf-8")
    text = _extract_text(f)
    assert "Åäö" in text


def test_latin1_with_replacement(tmp_path):
    """Files with invalid UTF-8 bytes should not crash — errors='replace'."""
    f = tmp_path / "latin.txt"
    f.write_bytes("Hej \xe5\xe4\xf6 d\xe4r.".encode("latin-1"))
    # Should not raise — errors="replace" is expected behaviour
    text = _extract_text(f)
    assert isinstance(text, str)
    assert len(text) > 0


def test_frontmatter_tags_block_style_unindented(tmp_path):
    # PyYAML's default dump writes list items without indentation
    f = tmp_path / "note.md"
    f.write_text("---\ntags:\n- pilot\n- lärare\n---\nText.", encoding="utf-8")
    assert _parse_md_frontmatter(f)["tags"] == ["pilot", "lärare"]


def test_frontmatter_tags_crlf(tmp_path):
    f = tmp_path / "note.md"
    f.write_bytes("---\r\ntitle: Win\r\ntags: [pilot, fas2]\r\n---\r\nText.".encode("utf-8"))
    fm = _parse_md_frontmatter(f)
    assert fm["title"] == "Win"
    assert fm["tags"] == ["pilot", "fas2"]


def test_frontmatter_tags_flow_comma_inside_quotes(tmp_path):
    f = tmp_path / "note.md"
    f.write_text('---\ntags: ["a, b", c]\n---\nText.', encoding="utf-8")
    assert _parse_md_frontmatter(f)["tags"] == ["a, b", "c"]


def test_frontmatter_tags_empty_and_duplicates_dropped(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\ntags: [pilot, , pilot, fas2,]\n---\nText.", encoding="utf-8")
    assert _parse_md_frontmatter(f)["tags"] == ["pilot", "fas2"]


def test_frontmatter_tags_scalar(tmp_path):
    f = tmp_path / "note.md"
    f.write_text("---\ntags: pilot\n---\nText.", encoding="utf-8")
    assert _parse_md_frontmatter(f)["tags"] == ["pilot"]


def test_md_frontmatter_stripped_crlf(tmp_path):
    f = tmp_path / "note.md"
    f.write_bytes("---\r\ntitle: Win\r\n---\r\nBrödtext.".encode("utf-8"))
    text = _extract_text(f)
    assert "title" not in text
    assert "Brödtext." in text


# ---------------------------------------------------------------------------
# .docx — tables and formatting offsets
# ---------------------------------------------------------------------------

def _make_docx(tmp_path, build):
    import docx
    d = docx.Document()
    build(d)
    f = tmp_path / "t.docx"
    d.save(str(f))
    return f


def test_docx_table_rows_imported_in_order(tmp_path):
    def build(d):
        d.add_paragraph("Intervju 1")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Intervjuare"
        t.cell(0, 1).text = "Hur började du?"
        t.cell(1, 0).text = "Anna"
        t.cell(1, 1).text = "Det var 2019."
        d.add_paragraph("Slut")
    text = _extract_text(_make_docx(tmp_path, build))
    assert text == "Intervju 1\nIntervjuare\tHur började du?\nAnna\tDet var 2019.\nSlut"


def test_docx_merged_cells_emitted_once(tmp_path):
    def build(d):
        t = d.add_table(rows=1, cols=3)
        a = t.cell(0, 0).merge(t.cell(0, 1))
        a.text = "Rubrik"
        t.cell(0, 2).text = "Text"
    assert _extract_text(_make_docx(tmp_path, build)) == "Rubrik\tText"


def test_docx_bold_span_offsets_match_text(tmp_path):
    from core.project import _extract_text_with_formatting
    def build(d):
        d.add_paragraph("Första")
        p = d.add_paragraph("Hej ")
        p.add_run("fet").bold = True
    text, spans = _extract_text_with_formatting(_make_docx(tmp_path, build))
    bold = [s for s in spans if s["type"] == "bold"]
    assert len(bold) == 1
    assert text[bold[0]["start"]:bold[0]["end"]] == "fet"
